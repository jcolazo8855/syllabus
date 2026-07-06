#!/usr/bin/env python3
"""Download the latest ABDC Journal Quality List and rebuild the local database.

Run from anywhere:  python research/abdc/scripts/update_jql.py

Outputs (under research/abdc/data/):
  abdc_jql_full.csv   — every journal on the list, all ratings, raw columns
  journals.json       — A*, A and B journals with derived author-guidelines links
  meta.json           — source URL, JQL version, counts, checksum, timestamp

The script discovers the current .xlsx link from the ABDC website; if the page
cannot be scraped it falls back to the last known direct URL recorded in
meta.json, then to a hardcoded default.
"""

import csv
import hashlib
import io
import json
import re
import sys
from datetime import datetime, timezone

import requests
from openpyxl import load_workbook

from common import (
    BROWSER_HEADERS,
    DATA_DIR,
    GUIDELINES_DIR,
    derive_guidelines,
    guidelines_cache_path,
    load_overrides,
    slugify,
)

LIST_PAGE = "https://abdc.edu.au/abdc-journal-quality-list/"
FALLBACK_XLSX = "https://abdc.edu.au/wp-content/uploads/2026/03/ABDC-JQL-2025-v1-260326.xlsx"
KEEP_RATINGS = {"A*", "A", "B"}


def discover_xlsx_url(session: requests.Session) -> str:
    """Find the newest JQL spreadsheet link on the ABDC list page."""
    candidates = []
    try:
        resp = session.get(LIST_PAGE, headers=BROWSER_HEADERS, timeout=60)
        resp.raise_for_status()
        links = re.findall(r'href="(https?://[^"]+\.xlsx?)"', resp.text, flags=re.I)
        # Ignore consultation drafts; prefer files that mention JQL
        for url in links:
            name = url.rsplit("/", 1)[-1].lower()
            if "draft" in name or "consultation" in name:
                continue
            if "jql" in name or "journal" in name:
                candidates.append(url)
    except requests.RequestException as exc:
        print(f"warning: could not scrape {LIST_PAGE}: {exc}", file=sys.stderr)

    if candidates:
        # wp-content upload paths embed /YYYY/MM/ — sort by that, newest first
        candidates.sort(reverse=True)
        return candidates[0]

    meta_path = DATA_DIR / "meta.json"
    if meta_path.exists():
        last = json.loads(meta_path.read_text()).get("source_xlsx")
        if last:
            print(f"falling back to previous source URL: {last}", file=sys.stderr)
            return last
    print(f"falling back to hardcoded URL: {FALLBACK_XLSX}", file=sys.stderr)
    return FALLBACK_XLSX


def find_header(rows: list[tuple]) -> tuple[int, dict]:
    """Locate the header row and map semantic fields to column indexes."""
    for i, row in enumerate(rows[:30]):
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        if any("journal" in c and "title" in c for c in cells) and any("rating" in c for c in cells):
            colmap = {}
            for j, c in enumerate(cells):
                if "journal" in c and "title" in c:
                    colmap["title"] = j
                elif "publisher" in c:
                    colmap["publisher"] = j
                elif "rating" in c:
                    colmap["rating"] = j
                elif "issn" in c and "online" in c:
                    colmap["issn_online"] = j
                elif c == "issn" or ("issn" in c and "online" not in c):
                    colmap.setdefault("issn", j)
                elif "field of research" in c:
                    colmap["field_of_research"] = j
                elif c.startswith("for") or "for code" in c:
                    colmap.setdefault("for_code", j)
                elif "inception" in c or "year" in c:
                    colmap["year_inception"] = j
            return i, colmap
    raise RuntimeError("could not locate header row in spreadsheet")


def clean(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"none", "nan"} else text


def main() -> int:
    session = requests.Session()
    url = discover_xlsx_url(session)
    print(f"downloading: {url}")
    resp = session.get(url, headers=BROWSER_HEADERS, timeout=120)
    resp.raise_for_status()
    blob = resp.content
    checksum = hashlib.sha256(blob).hexdigest()
    print(f"downloaded {len(blob):,} bytes  sha256={checksum[:16]}…")

    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    # The current list is usually the first sheet; older lists sit on later sheets.
    ws = wb.worksheets[0]
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    header_idx, colmap = find_header(rows)
    header = [clean(c) for c in rows[header_idx]]
    print(f"sheet '{ws.title}': header at row {header_idx + 1}: {header}")

    required = {"title", "rating"}
    if not required.issubset(colmap):
        raise RuntimeError(f"missing required columns, found only: {sorted(colmap)}")

    overrides = load_overrides()
    all_rows, journals = [], []
    seen_slugs = set()
    for row in rows[header_idx + 1:]:
        title = clean(row[colmap["title"]]) if colmap["title"] < len(row) else ""
        rating = clean(row[colmap["rating"]]) if colmap["rating"] < len(row) else ""
        if not title or not rating:
            continue
        record_raw = {header[j] or f"col{j}": clean(v) for j, v in enumerate(row) if j < len(header)}
        all_rows.append(record_raw)
        if rating not in KEEP_RATINGS:
            continue

        def col(name):
            j = colmap.get(name)
            return clean(row[j]) if j is not None and j < len(row) else ""

        slug = slugify(title)
        if slug in seen_slugs:  # rare duplicate titles — disambiguate by ISSN
            slug = slugify(f"{title}-{col('issn') or col('issn_online') or len(journals)}")
        seen_slugs.add(slug)

        publisher = col("publisher")
        issn, issn_online = col("issn"), col("issn_online")
        override = overrides.get(slug) or overrides.get(title)
        if override:
            guidelines = {"url": override["url"], "source": "override"}
        else:
            guidelines = derive_guidelines(title, slug, publisher, issn, issn_online)
        cache = guidelines_cache_path(slug)
        guidelines["cache"] = f"guidelines/{slug}.md" if cache.exists() else None

        journals.append(
            {
                "title": title,
                "slug": slug,
                "rating": rating,
                "publisher": publisher,
                "issn": issn,
                "issn_online": issn_online,
                "for_code": col("for_code"),
                "field_of_research": col("field_of_research"),
                "year_inception": col("year_inception"),
                "guidelines": guidelines,
            }
        )

    if len(journals) < 500:
        raise RuntimeError(
            f"parsed only {len(journals)} A*/A/B journals — spreadsheet format probably changed"
        )

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    GUIDELINES_DIR.mkdir(parents=True, exist_ok=True)

    fieldnames = list({k: None for rec in all_rows for k in rec})
    with open(DATA_DIR / "abdc_jql_full.csv", "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(all_rows)

    journals.sort(key=lambda r: (["A*", "A", "B"].index(r["rating"]), r["slug"]))
    version_match = re.search(r"(?:JQL[-_ ]?)(\d{4})", url, flags=re.I)
    counts = {r: sum(1 for j in journals if j["rating"] == r) for r in ["A*", "A", "B"]}
    payload = {
        "source": url,
        "jql_year": version_match.group(1) if version_match else None,
        "generated_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "counts": counts,
        "journals": journals,
    }
    (DATA_DIR / "journals.json").write_text(json.dumps(payload, indent=1, ensure_ascii=False))

    meta = {
        "source_xlsx": url,
        "sha256": checksum,
        "jql_year": payload["jql_year"],
        "generated_utc": payload["generated_utc"],
        "total_listed": len(all_rows),
        "kept": counts,
    }
    (DATA_DIR / "meta.json").write_text(json.dumps(meta, indent=2))
    print(f"done: {len(all_rows)} journals on list, kept {sum(counts.values())} ({counts})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
